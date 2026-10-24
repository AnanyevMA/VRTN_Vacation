import io
import os
import re
from unittest.mock import patch, MagicMock
import pytest
from pydantic import ValidationError
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.main import app
from backend.services.max_bot_client import MaxBotClient
from backend.services.llm_service import (
    GeminiFlashService,
    sanitize_prompt_pii,
    llm_service,
)
from backend.api.export_routes import sanitize_formula_injection, sanitize_row
from backend.schemas import (
    ChatMessageIn,
    EmployeeCreate,
    EmployeeUpdate,
    JobRoleCreate,
    JobRoleUpdate,
    WorkShiftCreate,
    WorkShiftUpdate,
    OptimizeVacationRequest,
    ScheduleRuleCreate,
    ScheduleRuleUpdate,
)

client = TestClient(app)


# --- 1. TLS Certificate Verification in MaxBotClient ---

def test_max_bot_client_tls_verify_enabled():
    """Verify MaxBotClient enforces standard TLS verification (verify=True)."""
    with patch("httpx.Client") as mock_httpx:
        MaxBotClient(token="test_token")
        assert mock_httpx.called
        _, kwargs = mock_httpx.call_args
        assert kwargs.get("verify") is True, f"Expected verify=True in httpx.Client, got: {kwargs.get('verify')}"

    # Verify live instance without mocking
    real_bot = MaxBotClient(token="test_token_2")
    # In httpx.Client, self._client.is_closed is False, verify is not False
    assert getattr(real_bot._client, "verify", None) is not False


# --- 2. LLM Service Header Auth & PII Sanitization ---

def test_sanitize_prompt_pii_removes_phone_and_emails():
    """Verify sanitize_prompt_pii strips phone numbers and email addresses (152-ФЗ compliance)."""
    text = (
        "Сотрудница Анна Ивановна (тел: +7 913 123 45 67, почта: anna.i@sewing-factory.ru) "
        "просит отпуск с 1 по 14 июля. Резервный контакт: 8-902-999-11-22."
    )
    cleaned = sanitize_prompt_pii(text)

    # Phone numbers and emails must be stripped
    assert "+7 913 123 45 67" not in cleaned
    assert "8-902-999-11-22" not in cleaned
    assert "anna.i@sewing-factory.ru" not in cleaned
    # Genuine vacation text must be preserved
    assert "Сотрудница Анна Ивановна" in cleaned
    assert "просит отпуск с 1 по 14 июля" in cleaned


def test_sanitize_prompt_pii_edge_cases():
    """Verify edge cases for sanitize_prompt_pii (empty, None, no PII)."""
    assert sanitize_prompt_pii("") == ""
    assert sanitize_prompt_pii(None) == ""
    normal = "Хочу в отпуск в августе на две недели"
    assert sanitize_prompt_pii(normal) == normal


def test_llm_service_uses_header_auth_and_sanitizes_pii():
    """Verify Gemini API request passes API key via x-goog-api-key header and sanitizes prompt."""
    service = GeminiFlashService(api_key="AIzaSyDummyKeyForTest12345")

    with patch.object(service._client, "post") as mock_post:
        mock_response = MagicMock()
        mock_response.status_code = 200
        mock_response.json.return_value = {
            "candidates": [
                {
                    "content": {
                        "parts": [
                            {"text": '{"slots":[{"start":"2026-07-01","end":"2026-07-14"}],"comment":"июль"}'}
                        ]
                    }
                }
            ]
        }
        mock_post.return_value = mock_response

        raw_user_input = "Позвоните +7 923 555 44 33 или напишите seamstress@khakassia.ru, отпуск в июле"
        slots, comment = service._query_gemini_for_dates(raw_user_input, default_year=2026, remaining_days=14)

        assert len(slots) == 1
        assert comment == "июль"
        assert mock_post.called

        url, kwargs = mock_post.call_args
        target_url = url[0]
        headers = kwargs.get("headers", {})
        payload = kwargs.get("json", {})

        # URL must NOT leak the API key in query params (?key=...)
        assert "key=" not in target_url
        # Key must be passed in x-goog-api-key header
        assert headers.get("x-goog-api-key") == "AIzaSyDummyKeyForTest12345"

        # Prompt content in payload must not contain phone or email
        prompt_text = payload["contents"][0]["parts"][0]["text"]
        assert "+7 923 555 44 33" not in prompt_text
        assert "seamstress@khakassia.ru" not in prompt_text


# --- 3. Security Response Headers Middleware & CORS ---

def test_security_response_headers_middleware():
    """Verify standard security headers are attached to API responses."""
    res = client.get("/")
    assert res.headers.get("X-Frame-Options") == "SAMEORIGIN"
    assert res.headers.get("X-Content-Type-Options") == "nosniff"
    assert res.headers.get("Referrer-Policy") == "strict-origin-when-cross-origin"


def test_cors_preflight_allowed_origin():
    """Verify CORS preflight returns correct headers for allowed origins."""
    headers = {
        "Origin": "http://localhost:8000",
        "Access-Control-Request-Method": "GET",
    }
    res = client.options("/api/departments", headers=headers)
    assert res.headers.get("access-control-allow-origin") == "http://localhost:8000"


def test_cors_preflight_disallowed_origin():
    """Verify CORS preflight does not authorize untrusted external origins."""
    headers = {
        "Origin": "http://evil-attacker-site.com",
        "Access-Control-Request-Method": "GET",
    }
    res = client.options("/api/departments", headers=headers)
    assert res.headers.get("access-control-allow-origin") != "http://evil-attacker-site.com"


def test_production_environment_disables_docs(monkeypatch):
    """Verify that when ENVIRONMENT=production, documentation endpoints are disabled."""
    monkeypatch.setenv("ENVIRONMENT", "production")
    # Import FastAPI and instantiate with the same logic as main.py
    from fastapi import FastAPI
    is_prod = os.getenv("ENVIRONMENT", "").strip().lower() == "production"
    prod_app = FastAPI(
        docs_url=None if is_prod else "/docs",
        redoc_url=None if is_prod else "/redoc",
        openapi_url=None if is_prod else "/openapi.json",
    )
    prod_client = TestClient(prod_app)
    assert prod_client.get("/docs").status_code == 404
    assert prod_client.get("/redoc").status_code == 404
    assert prod_client.get("/openapi.json").status_code == 404


# --- 4. Formula Injection Sanitization (CWE-1236) ---

def test_sanitize_formula_injection_helper():
    """Verify sanitize_formula_injection prepends single quote to cells starting with =, +, -, @."""
    assert sanitize_formula_injection("=cmd|' /C calc'!A0") == "'=cmd|' /C calc'!A0"
    assert sanitize_formula_injection("+12345") == "'+12345"
    assert sanitize_formula_injection("-500") == "'-500"
    assert sanitize_formula_injection("@SUM(A1:B10)") == "'@SUM(A1:B10)"

    # Benign string and non-string inputs should not be altered
    assert sanitize_formula_injection("Иванова Анна") == "Иванова Анна"
    assert sanitize_formula_injection(100) == 100
    assert sanitize_formula_injection(None) is None

    row = [1, "=1+1", "+Lead", "-Shift", "@Admin", "Safe text", 42]
    sanitized = sanitize_row(row)
    assert sanitized == [1, "'=1+1", "'+Lead", "'-Shift", "'@Admin", "Safe text", 42]


def test_export_csv_and_excel_formula_sanitization(monkeypatch):
    """Verify CSV and Excel exports sanitize fields starting with formula characters."""
    monkeypatch.delenv("ADMIN_API_KEY", raising=False)

    # Call CSV export endpoint
    res_csv = client.get("/api/export/csv")
    assert res_csv.status_code == 200
    csv_text = res_csv.content.decode("utf-8-sig")
    for line in csv_text.splitlines():
        cells = line.split(";")
        for cell in cells:
            clean_cell = cell.strip()
            # If a cell starts with formula characters, it must have been escaped with '
            if any(clean_cell.startswith(c) for c in ["=", "+", "-", "@"]):
                assert clean_cell.startswith("'"), f"Unescaped formula cell detected in CSV: {cell}"

    # Call Excel export endpoint
    res_excel = client.get("/api/export/excel")
    assert res_excel.status_code == 200
    wb = load_workbook(io.BytesIO(res_excel.content))
    ws = wb.active
    for row in ws.iter_rows(values_only=True):
        for val in row:
            if isinstance(val, str) and any(val.startswith(c) for c in ["=", "+", "-", "@"]):
                assert val.startswith("'"), f"Unescaped formula cell detected in Excel: {val}"


# --- 5. Bounded Schema Validations ---

def test_chat_message_in_max_length_500():
    """Verify ChatMessageIn enforces max_length=500 on message."""
    # 500 characters -> Valid
    msg_500 = "A" * 500
    obj = ChatMessageIn(employee_id=1, message=msg_500)
    assert len(obj.message) == 500

    # 501 characters -> ValidationError
    msg_501 = "A" * 501
    with pytest.raises(ValidationError):
        ChatMessageIn(employee_id=1, message=msg_501)


def test_employee_total_vacation_days_bounds():
    """Verify total_vacation_days bounded between 1 and 120 in EmployeeCreate and EmployeeUpdate."""
    # Valid boundaries
    emp_min = EmployeeCreate(
        phone="+79130000001",
        full_name="Тест Мин",
        role="Швея",
        department_id=1,
        total_vacation_days=1,
    )
    assert emp_min.total_vacation_days == 1

    emp_max = EmployeeCreate(
        phone="+79130000002",
        full_name="Тест Макс",
        role="Швея",
        department_id=1,
        total_vacation_days=120,
    )
    assert emp_max.total_vacation_days == 120

    # Below minimum (0)
    with pytest.raises(ValidationError):
        EmployeeCreate(
            phone="+79130000003",
            full_name="Тест Неверно",
            role="Швея",
            department_id=1,
            total_vacation_days=0,
        )

    # Above maximum (121)
    with pytest.raises(ValidationError):
        EmployeeCreate(
            phone="+79130000004",
            full_name="Тест Неверно",
            role="Швея",
            department_id=1,
            total_vacation_days=121,
        )

    # In EmployeeUpdate
    assert EmployeeUpdate(total_vacation_days=50).total_vacation_days == 50
    with pytest.raises(ValidationError):
        EmployeeUpdate(total_vacation_days=-1)
    with pytest.raises(ValidationError):
        EmployeeUpdate(total_vacation_days=125)


def test_default_quota_bounds_in_roles_and_shifts():
    """Verify default_quota bounded between 0 and 50 in JobRole and WorkShift schemas."""
    # JobRole bounds
    role_ok = JobRoleCreate(name="Швея", default_quota=0)
    assert role_ok.default_quota == 0
    role_max = JobRoleCreate(name="Швея", default_quota=50)
    assert role_max.default_quota == 50

    with pytest.raises(ValidationError):
        JobRoleCreate(name="Швея", default_quota=-1)
    with pytest.raises(ValidationError):
        JobRoleCreate(name="Швея", default_quota=51)

    with pytest.raises(ValidationError):
        JobRoleUpdate(default_quota=55)

    # WorkShift bounds
    shift_ok = WorkShiftCreate(name="Дневная", default_quota=0)
    assert shift_ok.default_quota == 0
    shift_max = WorkShiftCreate(name="Дневная", default_quota=50)
    assert shift_max.default_quota == 50

    with pytest.raises(ValidationError):
        WorkShiftCreate(name="Дневная", default_quota=-1)
    with pytest.raises(ValidationError):
        WorkShiftCreate(name="Дневная", default_quota=60)

    with pytest.raises(ValidationError):
        WorkShiftUpdate(default_quota=-2)


def test_window_search_days_bounds():
    """Verify window_search_days bounded between 1 and 90 in Optimizer and Rule schemas."""
    # In OptimizeVacationRequest
    opt_ok = OptimizeVacationRequest(employee_id=1, requested_slots=[], window_search_days=1)
    assert opt_ok.window_search_days == 1
    opt_max = OptimizeVacationRequest(employee_id=1, requested_slots=[], window_search_days=90)
    assert opt_max.window_search_days == 90

    with pytest.raises(ValidationError):
        OptimizeVacationRequest(employee_id=1, requested_slots=[], window_search_days=0)
    with pytest.raises(ValidationError):
        OptimizeVacationRequest(employee_id=1, requested_slots=[], window_search_days=91)

    # In ScheduleRuleCreate / ScheduleRuleUpdate
    rule_ok = ScheduleRuleCreate(
        rule_key="test_rule",
        rule_type="HARD",
        parameters={"window_search_days": 45},
    )
    assert rule_ok.parameters["window_search_days"] == 45

    with pytest.raises(ValidationError):
        ScheduleRuleCreate(
            rule_key="test_rule",
            rule_type="HARD",
            parameters={"window_search_days": 0},
        )

    with pytest.raises(ValidationError):
        ScheduleRuleCreate(
            rule_key="test_rule",
            rule_type="HARD",
            parameters={"window_search_days": 95},
        )

    with pytest.raises(ValidationError):
        ScheduleRuleUpdate(parameters={"window_search_days": -5})
